"""
Excel export service using openpyxl.
Generates branded, formatted Excel workbooks with multiple sheets.
"""

import io
from datetime import date

from openpyxl import Workbook
from openpyxl.styles import (
    Alignment,
    Border,
    Font,
    PatternFill,
    Side,
)
from openpyxl.utils import get_column_letter

from src.reports.schemas import (
    AttendanceReport,
    BatchReport,
    FeeCollectionReport,
    MonthlySummaryReport,
    PendingFeesReport,
    StudentPerformanceReport,
)

# ── Colours ─────────────────────────────────────────────────────────────────
HDR_FILL = PatternFill(start_color="1E1B4B", end_color="1E1B4B", fill_type="solid")
HDR_FONT = Font(color="FFFFFF", bold=True, size=10, name="Calibri")
KPI_FILL = PatternFill(start_color="EEF2FF", end_color="EEF2FF", fill_type="solid")
KPI_FONT = Font(color="4F46E5", bold=True, size=11, name="Calibri")
ALT_FILL = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
TITLE_FONT = Font(color="4F46E5", bold=True, size=14, name="Calibri")
NORMAL_FONT = Font(size=9, name="Calibri")
PAID_FONT = Font(color="10B981", bold=True, size=9, name="Calibri")
PARTIAL_FONT = Font(color="F59E0B", bold=True, size=9, name="Calibri")
PENDING_FONT = Font(color="EF4444", bold=True, size=9, name="Calibri")
THIN_BORDER = Border(
    left=Side(style="thin", color="E2E8F0"),
    right=Side(style="thin", color="E2E8F0"),
    top=Side(style="thin", color="E2E8F0"),
    bottom=Side(style="thin", color="E2E8F0"),
)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=False)
LEFT = Alignment(horizontal="left", vertical="center")


def _write_title(ws, title: str, subtitle: str, col_span: int) -> int:
    """Write title and subtitle, return next row index."""
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=col_span)
    ws.cell(1, 1, title).font = TITLE_FONT
    ws.cell(1, 1).alignment = LEFT
    ws.row_dimensions[1].height = 28

    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=col_span)
    sub_cell = ws.cell(2, 1, subtitle)
    sub_cell.font = Font(color="6B7280", size=9, name="Calibri")
    ws.row_dimensions[2].height = 18
    return 3  # next row


def _write_header_row(ws, row: int, headers: list[str]) -> None:
    for col, h in enumerate(headers, start=1):
        cell = ws.cell(row, col, h)
        cell.fill = HDR_FILL
        cell.font = HDR_FONT
        cell.alignment = CENTER
        cell.border = THIN_BORDER
    ws.row_dimensions[row].height = 22


def _write_data_row(ws, row: int, values: list, alt: bool = False, status: str | None = None) -> None:
    for col, v in enumerate(values, start=1):
        cell = ws.cell(row, col, v)
        cell.fill = ALT_FILL if alt else PatternFill()
        cell.border = THIN_BORDER
        cell.alignment = CENTER
        # Apply last column status font if applicable
    if status:
        last = ws.cell(row, len(values))
        last.font = {"paid": PAID_FONT, "partially_paid": PARTIAL_FONT, "pending": PENDING_FONT}.get(status, NORMAL_FONT)


def _auto_column_widths(ws) -> None:
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            if cell.value:
                max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[col_letter].width = min(max(max_len + 4, 10), 40)


def _currency(v: float) -> str:
    return f"₹{v:,.0f}"


def _pct(v: float) -> str:
    return f"{v:.1f}%"


# ══════════════════════════════════════════════════════════════════════════════
# Fee Collection Excel
# ══════════════════════════════════════════════════════════════════════════════

def generate_fee_collection_excel(report: FeeCollectionReport) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)

    # ── Summary Sheet ──
    ws_summary = wb.create_sheet("Summary")
    row = _write_title(ws_summary, "Fee Collection Report",
                        f"Period: {report.start_date} → {report.end_date}  |  Generated: {date.today()}", 4)

    kpi_items = [
        ("Total Expected", _currency(report.total_expected)),
        ("Total Collected", _currency(report.total_collected)),
        ("Total Pending", _currency(report.total_pending)),
        ("Collection Rate", _pct(report.collection_rate)),
        ("Paid Count", str(report.paid_count)),
        ("Partially Paid", str(report.partially_paid_count)),
        ("Pending Count", str(report.pending_count)),
        ("Total Students", str(report.total_students)),
    ]
    for label, value in kpi_items:
        lc = ws_summary.cell(row, 1, label)
        lc.font = Font(bold=True, size=9, name="Calibri")
        lc.border = THIN_BORDER
        vc = ws_summary.cell(row, 2, value)
        vc.fill = KPI_FILL
        vc.font = KPI_FONT
        vc.border = THIN_BORDER
        row += 1

    row += 1
    ws_summary.cell(row, 1, "Payment Mode Breakdown").font = Font(bold=True, size=10)
    row += 1
    _write_header_row(ws_summary, row, ["Mode", "Amount"])
    row += 1
    for mode, amt in report.by_mode.items():
        ws_summary.cell(row, 1, mode.replace("_", " ").title()).border = THIN_BORDER
        ws_summary.cell(row, 2, _currency(amt)).border = THIN_BORDER
        row += 1

    _auto_column_widths(ws_summary)

    # ── Detail Sheet ──
    ws_detail = wb.create_sheet("Fee Details")
    headers = ["Student", "Class", "Month", "Base Amt", "Discount", "Extra", "Net Amt", "Paid", "Balance", "Modes", "Last Payment", "Status"]
    row = _write_title(ws_detail, "Fee Collection — Detail", f"Period: {report.start_date} → {report.end_date}", len(headers))
    _write_header_row(ws_detail, row, headers)
    row += 1

    for i, r in enumerate(report.rows):
        values = [
            r.student_name, r.class_name, r.month,
            _currency(r.base_amount), _currency(r.discount), _currency(r.extra_charges),
            _currency(r.net_amount), _currency(r.paid_amount), _currency(r.balance),
            ", ".join(r.payment_modes),
            str(r.last_payment_date) if r.last_payment_date else "—",
            r.status.replace("_", " ").title(),
        ]
        _write_data_row(ws_detail, row, values, alt=(i % 2 == 1), status=r.status)
        row += 1

    _auto_column_widths(ws_detail)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ══════════════════════════════════════════════════════════════════════════════
# Pending Fees Excel
# ══════════════════════════════════════════════════════════════════════════════

def generate_pending_fees_excel(report: PendingFeesReport) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Pending Fees"

    headers = ["Student", "Class", "Month", "Net Amount", "Paid", "Balance", "Due Date", "Days Overdue", "Status"]
    row = _write_title(ws, "Pending Fees Report",
                        f"As of: {report.as_of_date}  |  Total Outstanding: {_currency(report.total_pending_amount)}", len(headers))
    _write_header_row(ws, row, headers)
    row += 1

    for i, r in enumerate(report.rows):
        values = [
            r.student_name, r.class_name, r.month,
            _currency(r.net_amount), _currency(r.paid_amount), _currency(r.balance),
            str(r.due_date), r.days_overdue, r.status.replace("_", " ").title(),
        ]
        _write_data_row(ws, row, values, alt=(i % 2 == 1), status=r.status)
        # Red background for severely overdue
        if r.days_overdue > 30:
            for col in range(1, len(headers) + 1):
                ws.cell(row, col).fill = PatternFill(start_color="FEF2F2", end_color="FEF2F2", fill_type="solid")
        row += 1

    _auto_column_widths(ws)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ══════════════════════════════════════════════════════════════════════════════
# Attendance Excel
# ══════════════════════════════════════════════════════════════════════════════

def generate_attendance_excel(report: AttendanceReport) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)

    # Student sheet
    ws_stu = wb.create_sheet("By Student")
    headers = ["Student", "Class", "Batch", "Sessions", "Present", "Absent", "Leave", "Attendance %"]
    row = _write_title(ws_stu, "Attendance Report — By Student",
                        f"Period: {report.start_date} → {report.end_date}  |  Overall Rate: {_pct(report.overall_rate)}", len(headers))
    _write_header_row(ws_stu, row, headers)
    row += 1
    for i, r in enumerate(report.student_rows):
        _write_data_row(ws_stu, row, [
            r.student_name, r.class_name, r.batch_name or "—",
            r.total_sessions, r.present, r.absent, r.leave,
            _pct(r.attendance_percentage),
        ], alt=(i % 2 == 1))
        pct_cell = ws_stu.cell(row, 8)
        pct_cell.font = PAID_FONT if r.attendance_percentage >= 75 else PENDING_FONT
        row += 1
    _auto_column_widths(ws_stu)

    # Daily sheet
    ws_day = wb.create_sheet("By Day")
    d_headers = ["Date", "Batch", "Total Students", "Present", "Absent", "Leave", "Rate"]
    row = _write_title(ws_day, "Attendance Report — Daily", f"Period: {report.start_date} → {report.end_date}", len(d_headers))
    _write_header_row(ws_day, row, d_headers)
    row += 1
    for i, r in enumerate(report.daily_rows):
        _write_data_row(ws_day, row, [
            str(r.date), r.batch_name or "All Batches",
            r.total_students, r.present, r.absent, r.leave,
            _pct(r.attendance_rate),
        ], alt=(i % 2 == 1))
        row += 1
    _auto_column_widths(ws_day)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ══════════════════════════════════════════════════════════════════════════════
# Student Performance Excel
# ══════════════════════════════════════════════════════════════════════════════

def generate_student_performance_excel(report: StudentPerformanceReport) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Student Performance"

    headers = ["Student", "Class", "Tests", "Avg Score %", "Highest %", "Lowest %", "Attendance %", "Outstanding Balance"]
    row = _write_title(ws, "Student Performance Report",
                        f"Period: {report.start_date} → {report.end_date}  |  Students: {len(report.rows)}", len(headers))
    _write_header_row(ws, row, headers)
    row += 1

    for i, r in enumerate(report.rows):
        _write_data_row(ws, row, [
            r.student_name, r.class_name, r.total_tests,
            _pct(r.avg_score_percentage), _pct(r.highest_score_percentage), _pct(r.lowest_score_percentage),
            _pct(r.attendance_percentage), _currency(r.outstanding_balance),
        ], alt=(i % 2 == 1))
        ws.cell(row, 4).font = PAID_FONT if r.avg_score_percentage >= 60 else PENDING_FONT
        ws.cell(row, 7).font = PAID_FONT if r.attendance_percentage >= 75 else PENDING_FONT
        ws.cell(row, 8).font = PENDING_FONT if r.outstanding_balance > 0 else PAID_FONT
        row += 1

    _auto_column_widths(ws)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ══════════════════════════════════════════════════════════════════════════════
# Batch Excel
# ══════════════════════════════════════════════════════════════════════════════

def generate_batch_excel(report: BatchReport) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Batch Report"

    headers = ["Batch", "Subject", "Teacher", "Days", "Timing", "Students", "Capacity", "Att. Rate", "Collected", "Pending"]
    row = _write_title(ws, "Batch Report",
                        f"Period: {report.start_date} → {report.end_date}  |  Batches: {len(report.rows)}", len(headers))
    _write_header_row(ws, row, headers)
    row += 1

    for i, r in enumerate(report.rows):
        _write_data_row(ws, row, [
            r.batch_name, r.subject, r.teacher, ", ".join(r.days[:3]), r.timing,
            r.total_students, r.max_students,
            _pct(r.attendance_rate), _currency(r.total_fee_collected), _currency(r.total_fee_pending),
        ], alt=(i % 2 == 1))
        row += 1

    _auto_column_widths(ws)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ══════════════════════════════════════════════════════════════════════════════
# Monthly Summary Excel
# ══════════════════════════════════════════════════════════════════════════════

def generate_monthly_summary_excel(report: MonthlySummaryReport) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Monthly Summary"

    headers = ["Month", "Active Students", "Sessions", "Avg Att. %", "Expected", "Collected", "Pending", "Coll. Rate", "Tests"]
    row = _write_title(ws, "Monthly Summary Report",
                        f"Period: {report.start_date} → {report.end_date}", len(headers))
    _write_header_row(ws, row, headers)
    row += 1

    for i, r in enumerate(report.rows):
        _write_data_row(ws, row, [
            r.month, r.active_students, r.total_sessions,
            _pct(r.avg_attendance_rate),
            _currency(r.fees_expected), _currency(r.fees_collected), _currency(r.fees_pending),
            _pct(r.collection_rate), r.tests_conducted,
        ], alt=(i % 2 == 1))
        cr_cell = ws.cell(row, 8)
        if r.collection_rate >= 80:
            cr_cell.font = PAID_FONT
        elif r.collection_rate >= 50:
            cr_cell.font = PARTIAL_FONT
        else:
            cr_cell.font = PENDING_FONT
        row += 1

    _auto_column_widths(ws)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
