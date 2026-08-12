"""
ExportService — generates CSV and Excel downloads for student data.
"""

from __future__ import annotations

import csv
import io
from collections.abc import Sequence

from sqlalchemy import or_, select
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy.orm import selectinload

from src.students.models import Batch, Student


EXPORT_COLUMNS = [
    "name",
    "parent_name",
    "parent_mobile",
    "alternate_mobile",
    "class_name",
    "school",
    "subjects",
    "joining_date",
    "monthly_fee",
    "address",
    "notes",
    "is_active",
    "is_demo",
    "batch_names",
]

COLUMN_LABELS = {
    "name": "Student Name",
    "parent_name": "Parent Name",
    "parent_mobile": "Parent Mobile",
    "alternate_mobile": "Alternate Mobile",
    "class_name": "Class",
    "school": "School",
    "subjects": "Subjects",
    "joining_date": "Joining Date",
    "monthly_fee": "Monthly Fee",
    "address": "Address",
    "notes": "Notes",
    "is_active": "Active",
    "is_demo": "Demo Record",
    "batch_names": "Batches",
}


class ExportService:
    """Generates CSV and Excel exports of student data with filtering support."""

    async def get_students(
        self,
        db: AsyncSession,
        status: str = "all",  # all | active | inactive
        batch_id: str | None = None,
        class_name: str | None = None,
        search: str | None = None,
    ) -> Sequence[Student]:
        """Fetch students from DB with filters applied."""
        stmt = select(Student).options(selectinload(Student.batches))

        if status == "active":
            stmt = stmt.where(Student.is_active == True)  # noqa: E712
        elif status == "inactive":
            stmt = stmt.where(Student.is_active == False)  # noqa: E712

        if batch_id:
            stmt = stmt.join(Student.batches).where(Batch.id == batch_id)

        if class_name:
            stmt = stmt.where(Student.class_name == class_name)

        if search:
            q = f"%{search.strip()}%"
            stmt = stmt.where(
                or_(
                    Student.name.ilike(q),
                    Student.parent_name.ilike(q),
                    Student.parent_mobile.ilike(q),
                    Student.school.ilike(q),
                    Student.class_name.ilike(q),
                )
            )

        stmt = stmt.order_by(Student.name)
        result = await db.execute(stmt)
        return result.scalars().all()

    def _student_to_row(self, student: Student) -> dict:
        """Convert a Student ORM object to an export row dict."""
        return {
            "name": student.name,
            "parent_name": student.parent_name,
            "parent_mobile": student.parent_mobile,
            "alternate_mobile": student.alternate_mobile or "",
            "class_name": student.class_name,
            "school": student.school,
            "subjects": ", ".join(student.subjects) if student.subjects else "",
            "joining_date": student.joining_date.strftime("%Y-%m-%d") if student.joining_date else "",
            "monthly_fee": str(student.monthly_fee),
            "address": student.address,
            "notes": student.notes or "",
            "is_active": "Yes" if student.is_active else "No",
            "is_demo": "Yes" if student.is_demo else "No",
            "batch_names": ", ".join(b.name for b in student.batches) if student.batches else "",
        }

    def export_csv(self, students: Sequence[Student]) -> bytes:
        """Generate a UTF-8 CSV file from a list of students."""
        output = io.StringIO()
        writer = csv.DictWriter(output, fieldnames=EXPORT_COLUMNS)

        # Write header row with human-readable labels
        writer.writerow({col: COLUMN_LABELS[col] for col in EXPORT_COLUMNS})

        for student in students:
            writer.writerow(self._student_to_row(student))

        return output.getvalue().encode("utf-8-sig")  # BOM for Excel compatibility

    def export_excel(self, students: Sequence[Student]) -> bytes:
        """Generate a formatted .xlsx file from a list of students."""
        import openpyxl
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Students"

        # ── Header row styling ──
        header_fill = PatternFill(start_color="4F46E5", end_color="4F46E5", fill_type="solid")
        header_font = Font(color="FFFFFF", bold=True, size=11)
        header_alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        for col_idx, col_key in enumerate(EXPORT_COLUMNS, start=1):
            cell = ws.cell(row=1, column=col_idx, value=COLUMN_LABELS[col_key])
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = header_alignment

        ws.row_dimensions[1].height = 30

        # ── Data rows ──
        alt_fill = PatternFill(start_color="EEF2FF", end_color="EEF2FF", fill_type="solid")
        for row_idx, student in enumerate(students, start=2):
            row_data = self._student_to_row(student)
            fill = alt_fill if row_idx % 2 == 0 else None
            for col_idx, col_key in enumerate(EXPORT_COLUMNS, start=1):
                cell = ws.cell(row=row_idx, column=col_idx, value=row_data[col_key])
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                if fill:
                    cell.fill = fill

        # ── Column widths ──
        col_widths = {
            "name": 22,
            "parent_name": 22,
            "parent_mobile": 16,
            "alternate_mobile": 16,
            "class_name": 10,
            "school": 28,
            "subjects": 22,
            "joining_date": 14,
            "monthly_fee": 14,
            "address": 36,
            "notes": 30,
            "is_active": 10,
            "is_demo": 12,
            "batch_names": 28,
        }
        for col_idx, col_key in enumerate(EXPORT_COLUMNS, start=1):
            ws.column_dimensions[get_column_letter(col_idx)].width = col_widths.get(col_key, 20)

        # Freeze the header row
        ws.freeze_panes = "A2"

        output = io.BytesIO()
        wb.save(output)
        return output.getvalue()
